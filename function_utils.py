import csv
import re
import html
import uuid
import os
import json
import urllib.parse
import urllib.request
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

BASE_DIR = Path(__file__).resolve().parent
BOOK_DATA_PATH = BASE_DIR / "BOOKS_DATA.xlsx"
STORE_DATA_PATH = BASE_DIR / "STORE_RECORDS.xlsx"
EMPLS_DATA_PATH = BASE_DIR / "EMPLOYEES.csv"
LOG_DATA_PATH = BASE_DIR / "Activity_Log.txt"
GENRES_DATA_PATH = BASE_DIR / "GENRES.txt"
EXPENSES_DATA_PATH = BASE_DIR / "EXPENSES.csv"
CREDT_DATA_PATH = BASE_DIR / "CREDENTIAL.txt"
ADDRESS_CACHE_PATH = BASE_DIR / "ADDRESS_CACHE.json"
GEMINI_MODEL = "gemini-3.5-flash-lite"
DESIGNATION_ACCESS = {
    "store clerk": "s",
    "inventory manager": "s",
    "sales executive": "s",
    "assistant manager": "s",
    "director": "csp",
}
DEFAULT_EMPLOYEE = {"Employee ID": "EMP101", "Full Name": "Default Director", "Phone": "", "Email": "", "Username": "user@", "Password": "12345678", "Designation": "Director", "Access": "csp", "Status": "Active", "Joined On": datetime.now().strftime("%d-%m-%Y")}

BOOK_HEADERS = ["Book ID", "Book Name", "Author Name", "Genre", "Language", "Published Date", "Wholesale Price", "Market Price", "Profit Margin", "Quantities Available"]
EMP_HEADERS = ["Employee ID", "Full Name", "Phone", "Email", "Username", "Password", "Designation", "Access", "Status", "Joined On"]
SALES_HEADERS = ["Order ID", "Order Date", "Customer Name", "Phone", "Flat / House / Building", "Street / Area", "Landmark", "City", "State", "PIN", "Payment Method", "Book ID", "Book Name", "Author", "Genre", "Quantity", "Unit Price", "Cost Price", "Line Total", "Line Profit", "Order Status"]
EXPENSE_HEADERS = ["Expense ID", "Date", "Category", "Description", "Amount"]

ALIASES = {
    "id": ["book id", "book code", "ddc code", "code", "id", "isbn"],
    "name": ["book name", "title", "name", "book"],
    "author": ["author name", "author", "writer"],
    "genre": ["genre", "category", "section"],
    "language": ["language", "lang"],
    "published": ["published date", "publication date", "year", "published"],
    "cost": ["wholesale price", "cost price", "purchase price", "cost"],
    "price": ["market price (inr)", "market price", "selling price", "sale price", "price", "mrp"],
    "profit": ["profit margin", "margin", "profit"],
    "stock": ["quantities available", "quantity", "stock", "copies", "available copies"]
}

def norm(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()

def clean_name(value):
    return " ".join(str(value or "").strip().split()).title()

def clean_address(value):
    value = " ".join(str(value or "").strip().split())
    if not value:
        return ""
    parts = re.split(r",\s*", value)
    return ", ".join(p.strip().title() for p in parts if p.strip())

def clean_text(value):
    return " ".join(str(value or "").strip().split())

def money(value):
    try:
        return round(float(value or 0), 2)
    except Exception:
        return 0.0

def integer(value):
    try:
        return int(float(value or 0))
    except Exception:
        return 0

def _seed_default_employee():
    try:
        if not EMPLS_DATA_PATH.exists():
            return
        with open(EMPLS_DATA_PATH, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            rows = [dict(row) for row in reader]
        matched = False
        for row in rows:
            if clean_text(row.get("Username", "")).lower() == DEFAULT_EMPLOYEE["Username"].lower():
                row.update(DEFAULT_EMPLOYEE)
                matched = True
                break
        if not matched:
            rows.append(dict(DEFAULT_EMPLOYEE))
        with open(EMPLS_DATA_PATH, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=EMP_HEADERS)
            writer.writeheader()
            writer.writerows({header: row.get(header, "") for header in EMP_HEADERS} for row in rows)
    except Exception:
        pass

def ensure_runtime_files():
    BASE_DIR.mkdir(parents=True, exist_ok=True)
    if not BOOK_DATA_PATH.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "BOOKS -- GENERAL"
        ws.append(BOOK_HEADERS)
        wb.save(BOOK_DATA_PATH)
    if not STORE_DATA_PATH.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(SALES_HEADERS)
        wb.create_sheet("Orders")
        wb.save(STORE_DATA_PATH)
    if not EMPLS_DATA_PATH.exists():
        with open(EMPLS_DATA_PATH, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerow(EMP_HEADERS)
    _seed_default_employee()
    if not LOG_DATA_PATH.exists():
        LOG_DATA_PATH.write_text("BOOK STORE ACTIVITY LOG\n" + "=" * 90 + "\n", encoding="utf-8")
    if not GENRES_DATA_PATH.exists():
        GENRES_DATA_PATH.write_text("General\n", encoding="utf-8")
    if not EXPENSES_DATA_PATH.exists():
        with open(EXPENSES_DATA_PATH, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerow(EXPENSE_HEADERS)
    if not CREDT_DATA_PATH.exists():
        CREDT_DATA_PATH.write_text("Runtime credentials are managed through EMPLOYEES.csv.\n", encoding="utf-8")
    if not ADDRESS_CACHE_PATH.exists():
        ADDRESS_CACHE_PATH.write_text("{}", encoding="utf-8")


def _gemini_api_key():
    key = os.getenv("GEMINI_API_KEY", "").strip()
    if key:
        return key
    try:
        import streamlit as st
        return str(st.secrets.get("GEMINI_API_KEY", "")).strip()
    except Exception:
        return ""

def _address_cache():
    ensure_runtime_files()
    try:
        data = json.loads(ADDRESS_CACHE_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}

def _save_address_cache(data):
    ADDRESS_CACHE_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

def _gemini_json(prompt, cache_key):
    cache = _address_cache()
    if cache_key in cache and isinstance(cache[cache_key], list) and cache[cache_key]:
        return cache[cache_key]
    api_key = _gemini_api_key()
    if not api_key:
        return []
    endpoint = f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent?key={urllib.parse.quote(api_key)}"
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": 0,
            "responseMimeType": "application/json",
            "maxOutputTokens": 8192,
            "thinkingConfig": {"thinkingLevel": "minimal"}
        }
    }
    try:
        req = urllib.request.Request(endpoint, data=json.dumps(payload).encode("utf-8"), headers={"Content-Type":"application/json"}, method="POST")
        with urllib.request.urlopen(req, timeout=18) as response:
            raw=json.loads(response.read().decode("utf-8"))
        text=""
        for candidate in raw.get("candidates", []):
            for part in candidate.get("content", {}).get("parts", []):
                if part.get("text"):
                    text += part["text"]
        data=json.loads(text.strip())
        if isinstance(data, dict):
            data=data.get("items", data.get("values", []))
        if not isinstance(data, list):
            return []
        result=sorted({str(x).strip() for x in data if str(x).strip()}, key=str.casefold)
        if result:
            cache[cache_key]=result
            _save_address_cache(cache)
        return result
    except Exception:
        return []

def get_indian_states():
    return _gemini_json("Return a JSON array containing every current Indian state and union territory name. Output only the array.", "states")

def get_indian_districts(state):
    if not state or state == "Address list unavailable":
        return []
    key=f"districts::{state}"
    return _gemini_json(f"Return a JSON array containing every current district of the Indian state or union territory '{state}'. Use official district names. Output only the array.", key)

def get_indian_cities(state, district):
    if not state or not district or "unavailable" in {state.lower(), district.lower()}:
        return []
    key=f"cities::{state}::{district}"
    return _gemini_json(f"Return a JSON array containing cities, towns and officially recognized urban localities in district '{district}', '{state}', India. Use names suitable for an address selector. Output only the array.", key)

def get_city_pincodes(state, district, city):
    if not state or not district or not city or "unavailable" in {state.lower(), district.lower(), city.lower()}:
        return []
    key=f"pincodes::{state}::{district}::{city}"
    values=_gemini_json(f"Return a JSON array of all valid six-digit India Post PIN codes serving '{city}', district '{district}', '{state}', India. Do not invent or estimate PIN codes. Output only the six-digit codes.", key)
    return sorted({str(x).zfill(6) for x in values if str(x).isdigit() and len(str(x)) <= 6})

def _header_map(headers):
    return {norm(h): i for i, h in enumerate(headers)}

def _find_index(mapping, aliases):
    for alias in aliases:
        if norm(alias) in mapping:
            return mapping[norm(alias)]
    for key, idx in mapping.items():
        for alias in aliases:
            if norm(alias) in key or key in norm(alias):
                return idx
    return None

def load_books():
    ensure_runtime_files()
    wb = load_workbook(BOOK_DATA_PATH, data_only=True)
    books = []
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        if not values:
            continue
        headers = list(values[0])
        mapping = _header_map(headers)
        for row in values[1:]:
            if not row or not any(v is not None and str(v).strip() for v in row):
                continue
            def get(kind):
                idx = _find_index(mapping, ALIASES[kind])
                return row[idx] if idx is not None and idx < len(row) else ""
            bid = clean_text(get("id"))
            name = clean_text(get("name"))
            if not name:
                continue
            genre = clean_text(get("genre")) or ws.title.split("--")[0].strip().title()
            books.append({
                "id": bid,
                "name": name,
                "author": clean_name(get("author")),
                "genre": genre,
                "language": clean_text(get("language")) or "English",
                "published": clean_text(get("published")),
                "cost": money(get("cost")),
                "price": money(get("price")),
                "profit": money(get("profit")),
                "stock": integer(get("stock")),
                "sheet": ws.title
            })
    used = {b["id"] for b in books if b["id"]}
    seq = 1
    for b in books:
        if not b["id"]:
            while f"BK{seq:04d}" in used:
                seq += 1
            b["id"] = f"BK{seq:04d}"
            used.add(b["id"])
            seq += 1
    return books

def save_books(books):
    wb = Workbook()
    first = True
    groups = {}
    for b in books:
        groups.setdefault(b.get("genre") or "General", []).append(b)
    for genre, rows in groups.items():
        title = re.sub(r"[^A-Za-z0-9 &_-]", "", genre)[:20] or "General"
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title[:31]
        ws.append(BOOK_HEADERS)
        for i, b in enumerate(rows, 1):
            ws.append([b["id"], b["name"], b["author"], b["genre"], b["language"], b["published"], b["cost"], b["price"], b["profit"], b["stock"]])
    if first:
        ws = wb.active
        ws.title = "BOOKS -- GENERAL"
        ws.append(BOOK_HEADERS)
    wb.save(BOOK_DATA_PATH)

def get_book(book_id):
    for b in load_books():
        if b["id"] == book_id:
            return b
    return None

def add_book(book):
    books = load_books()
    books.append(book)
    save_books(books)
    return book

def update_book(book_id, **changes):
    books = load_books()
    found = None
    for b in books:
        if b["id"] == book_id:
            b.update(changes)
            found = b
            break
    if found:
        save_books(books)
    return found

def add_stock(book_id, copies):
    books = load_books()
    for b in books:
        if b["id"] == book_id:
            b["stock"] += integer(copies)
            save_books(books)
            return b
    return None

def reduce_stock(items):
    books = load_books()
    index = {b["id"]: b for b in books}
    for item in items:
        if item["book_id"] not in index or index[item["book_id"]]["stock"] < integer(item["quantity"]):
            return False
    for item in items:
        index[item["book_id"]]["stock"] -= integer(item["quantity"])
    save_books(books)
    return True

def load_employees():
    ensure_runtime_files()
    _seed_default_employee()
    with open(EMPLS_DATA_PATH, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if not rows:
        return []
    headers = rows[0]
    result = []
    for row in rows[1:]:
        if not row or not any(str(x).strip() for x in row):
            continue
        row = row + [""] * max(0, len(EMP_HEADERS) - len(row))
        if len(headers) < len(EMP_HEADERS):
            headers = EMP_HEADERS
        result.append(dict(zip(headers, row)))
    return result

def generate_employee_id():
    highest = 100
    for emp in load_employees():
        text = " ".join(str(v) for v in emp.values())
        for n in re.findall(r"EMP\s*0*(\d+)", text.upper()):
            highest = max(highest, int(n))
    return f"EMP{highest + 1}"

def authenticate_employee(username, password, requested_portal=None):
    username = clean_text(username)
    password = str(password)
    for emp in load_employees():
        if clean_text(emp.get("Username")) == username and str(emp.get("Password", "")) == password:
            access = clean_text(emp.get("Access")).lower()
            designation = clean_text(emp.get("Designation")).lower()
            if username.lower() == DEFAULT_EMPLOYEE["Username"].lower() and password == DEFAULT_EMPLOYEE["Password"]:
                role = requested_portal if requested_portal in {"Store Clerk", "Director"} else "Director"
            else:
                role = "Director" if access in {"csp", "p", "director", "admin"} or "director" in designation else "Store Clerk"
            emp["role"] = role
            return emp
    return None

def create_employee(full_name, phone, email, username, password, designation="Store Clerk"):
    employees = load_employees()
    if any(clean_text(e.get("Username")).lower() == clean_text(username).lower() for e in employees):
        raise ValueError("Username already exists.")
    emp_id = generate_employee_id()
    designation = clean_name(designation)
    access = DESIGNATION_ACCESS.get(designation.lower(), "s")
    record = {"Employee ID": emp_id, "Full Name": clean_name(full_name), "Phone": clean_text(phone), "Email": clean_text(email).lower(), "Username": clean_text(username), "Password": str(password), "Designation": designation, "Access": access, "Status": "Active", "Joined On": datetime.now().strftime("%d-%m-%Y")}
    with open(EMPLS_DATA_PATH, "a", newline="", encoding="utf-8") as f:
        csv.writer(f).writerow([record[h] for h in EMP_HEADERS])
    return record

def load_genres():
    ensure_runtime_files()
    values = [clean_text(x) for x in GENRES_DATA_PATH.read_text(encoding="utf-8").splitlines() if clean_text(x)]
    values += [b["genre"] for b in load_books() if b["genre"]]
    return sorted(set(values), key=str.lower)

def add_genre(genre):
    genre = clean_name(genre)
    if not genre:
        return False
    genres = load_genres()
    if genre.lower() in {g.lower() for g in genres}:
        return False
    with open(GENRES_DATA_PATH, "a", encoding="utf-8") as f:
        f.write(genre + "\n")
    return True

def log_activity(actor, action, details=""):
    ensure_runtime_files()
    stamp = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
    with open(LOG_DATA_PATH, "a", encoding="utf-8") as f:
        f.write(f"[{stamp}] | {clean_text(actor)} | {clean_text(action)} | {clean_text(details)}\n")

def load_logs(limit=250):
    ensure_runtime_files()
    lines = LOG_DATA_PATH.read_text(encoding="utf-8").splitlines()
    return [x for x in lines if x.startswith("[")][-limit:][::-1]

def _ensure_sales_sheet(wb):
    if "Sales" not in wb.sheetnames:
        ws = wb.create_sheet("Sales")
        ws.append(SALES_HEADERS)
    return wb["Sales"]

def append_order(order, items):
    ensure_runtime_files()
    wb = load_workbook(STORE_DATA_PATH)
    ws = _ensure_sales_sheet(wb)
    for item in items:
        ws.append([order["order_id"], order["date"], order["customer_name"], order["phone"], order["flat"], order["street"], order["landmark"], order["city"], order["state"], order["pin"], order["payment"], item["book_id"], item["name"], item["author"], item["genre"], item["quantity"], item["unit_price"], item["cost_price"], item["line_total"], item["line_profit"], "Confirmed"])
    if "Orders" not in wb.sheetnames:
        wb.create_sheet("Orders")
    ow = wb["Orders"]
    if ow.max_row == 1 and ow.cell(1,1).value is None:
        ow.append(["Order ID", "Date", "Customer Name", "Phone", "City", "State", "PIN", "Payment", "Items", "Grand Total", "Grand Profit", "Status"])
    elif ow.max_row == 1 and ow.cell(1,1).value != "Order ID":
        pass
    if ow.max_row == 1 and ow.cell(1,1).value is None:
        pass
    if ow.max_row == 1 and ow.cell(1,1).value == "Order ID":
        pass
    ow.append([order["order_id"], order["date"], order["customer_name"], order["phone"], order["city"], order["state"], order["pin"], order["payment"], sum(integer(i["quantity"]) for i in items), order["grand_total"], order["grand_profit"], "Confirmed"])
    wb.save(STORE_DATA_PATH)

def new_order_id():
    return "ORD-" + datetime.now().strftime("%Y%m%d") + "-" + uuid.uuid4().hex[:6].upper()

def load_sales():
    ensure_runtime_files()
    wb = load_workbook(STORE_DATA_PATH, data_only=True)
    if "Sales" not in wb.sheetnames:
        return []
    ws = wb["Sales"]
    values = list(ws.values)
    if not values:
        return []
    headers = [str(x) for x in values[0]]
    return [dict(zip(headers, r)) for r in values[1:] if r and any(x is not None for x in r)]

def add_expense(category, description, amount):
    eid = "EXP-" + datetime.now().strftime("%Y%m%d%H%M%S")
    with open(EXPENSES_DATA_PATH, "a", newline="", encoding="utf-8") as f:
        csv.writer(f).writerow([eid, datetime.now().strftime("%d-%m-%Y"), clean_name(category), clean_text(description), money(amount)])
    return eid

def load_expenses():
    ensure_runtime_files()
    with open(EXPENSES_DATA_PATH, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    return rows

def invoice_html(order, items):
    rows = "".join(
        f"<tr><td><b>📚</b> {html.escape(str(i['name']))}<br><span class='muted'>{html.escape(str(i.get('author','')))}</span></td>"
        f"<td>{html.escape(str(i['book_id']))}</td><td>{i['quantity']}</td>"
        f"<td>₹{i['unit_price']:.2f}</td><td><b>₹{i['line_total']:.2f}</b></td></tr>"
        for i in items
    )
    address = ", ".join(
        x for x in [
            order.get("flat", ""),
            order.get("street", ""),
            order.get("landmark", ""),
            order.get("city", ""),
            f"{order.get('state','')} - {order.get('pin','')}" if order.get("state") else ""
        ] if x
    )
    return f'''<!doctype html>
<html><head><meta charset="utf-8"><title>{html.escape(str(order['order_id']))}</title>
<style>
body{{font-family:Arial,sans-serif;background:linear-gradient(135deg,#eef2f7,#f7f3ff);color:#172033;padding:30px}}
.invoice{{max-width:980px;margin:auto;background:#fff;border:1px solid #e4e7ec;border-radius:24px;overflow:hidden;box-shadow:0 18px 55px rgba(23,32,51,.12)}}
.header{{padding:30px 34px;background:linear-gradient(120deg,#17152f,#3c327e);color:#fff;display:flex;justify-content:space-between;gap:24px}}
.brand{{font-size:30px;font-weight:800;letter-spacing:.5px}} .sub{{opacity:.78;margin-top:5px}}
.orderbox{{background:rgba(255,255,255,.12);border:1px solid rgba(255,255,255,.22);border-radius:15px;padding:14px 18px;min-width:190px}}
.content{{padding:30px 34px}} .grid{{display:grid;grid-template-columns:1fr 1fr;gap:16px}}
.panel{{background:#f7f8fb;border:1px solid #e7eaf0;border-radius:16px;padding:17px}}
.panel h3{{margin:0 0 8px;color:#342a79}} .muted{{color:#667085;font-size:13px}}
table{{width:100%;border-collapse:separate;border-spacing:0;margin-top:22px;border:1px solid #e4e7ec;border-radius:15px;overflow:hidden}}
th,td{{padding:13px 12px;border-bottom:1px solid #e9ebef;text-align:left}} th{{background:#eeeaff;color:#342a79;font-size:13px}} tr:last-child td{{border-bottom:0}}
.summary{{margin-top:22px;display:flex;justify-content:flex-end}} .summarybox{{width:340px;background:#f7f8fb;border:1px solid #e4e7ec;border-radius:16px;padding:18px}}
.total{{font-size:25px;font-weight:800;color:#087f70;margin-top:9px}}
.pay{{margin-top:18px;padding:15px 17px;background:#ecfdf5;border:1px solid #b7ebd2;border-radius:14px;color:#075e49}}
.footer{{background:#11132b;color:#fff;padding:25px 34px;text-align:center}} .footer-title{{font-size:19px;font-weight:800}} .footer-sub{{font-size:12px;opacity:.82;margin-top:5px}} .footer-copy{{font-size:10px;opacity:.62;margin-top:13px;line-height:1.7}}
</style></head>
<body><div class="invoice">
<div class="header"><div><div class="brand">📚 BOOKSKART</div><div class="sub">Sales Invoice · Order Fulfilment Record</div></div>
<div class="orderbox"><b>ORDER ID</b><br>{html.escape(str(order['order_id']))}<br><span style="opacity:.75">{html.escape(str(order['date']))}</span></div></div>
<div class="content">
<div class="grid"><div class="panel"><h3>Customer</h3><b>{html.escape(str(order['customer_name']))}</b><br>{html.escape(str(order['phone']))}</div>
<div class="panel"><h3>Delivery Address</h3>{html.escape(address)}</div></div>
<table><tr><th>Book</th><th>Book Code</th><th>Qty</th><th>Unit Price</th><th>Line Total</th></tr>{rows}</table>
<div class="summary"><div class="summarybox"><div class="muted">Order Summary</div><div style="margin-top:8px">Items: <b>{sum(int(i['quantity']) for i in items)}</b></div>
<div style="margin-top:5px">Payment: <b>{html.escape(str(order['payment']))}</b></div><div class="total">Grand Total: ₹{order['grand_total']:,.2f}</div></div></div>
<div class="pay">✓ Order Status: <b>Confirmed</b> · Payment Method: <b>{html.escape(str(order['payment']))}</b></div>
</div>
<div class="footer"><div class="footer-title">BooksKart</div>
<div class="footer-sub">Crafted For Readers, Designed By MJJ-TechWorld</div>
<div class="footer-sub">Your Data Is Safe &amp; Private • Customer information is used for order processing and store records.</div>
<div class="footer-copy">© 2026 BooksKart. All Rights Reserved. | Terms Of Service | Privacy Policy | Disclaimer<br>
Designed &amp; Developed By MJJ-TechWorld • Made In India • Support: support@mjjtechworld.com • Version 2.0</div></div>
</div></body></html>'''


def invoice_pdf(order, items):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, rightMargin=32, leftMargin=32, topMargin=34, bottomMargin=48
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "InvoiceTitle", parent=styles["Title"], fontSize=23, leading=27,
        textColor=colors.HexColor("#ffffff"), spaceAfter=4
    )
    white_small = ParagraphStyle(
        "WhiteSmall", parent=styles["Normal"], fontSize=9.5, leading=12,
        textColor=colors.HexColor("#dfe4ff")
    )
    section = ParagraphStyle(
        "Section", parent=styles["Heading3"], fontSize=11, leading=14,
        textColor=colors.HexColor("#342a79"), spaceAfter=6
    )
    normal = ParagraphStyle(
        "InvoiceNormal", parent=styles["Normal"], fontSize=9.2, leading=13,
        textColor=colors.HexColor("#172033")
    )
    small = ParagraphStyle(
        "InvoiceSmall", parent=styles["Normal"], fontSize=8.2, leading=11,
        textColor=colors.HexColor("#667085")
    )
    total_style = ParagraphStyle(
        "InvoiceTotal", parent=normal, alignment=TA_RIGHT, fontSize=16,
        leading=19, textColor=colors.HexColor("#087f70")
    )
    footer_style = ParagraphStyle(
        "InvoiceFooter", parent=styles["Normal"], fontSize=7.5, leading=10,
        alignment=1, textColor=colors.HexColor("#69718c")
    )

    address = ", ".join(
        x for x in [
            order.get("flat", ""), order.get("street", ""), order.get("landmark", ""),
            order.get("city", ""),
            f"{order.get('state','')} - {order.get('pin','')}" if order.get("state") else ""
        ] if x
    )

    story = []
    header = Table([[
        Paragraph("BOOKSKART", title),
        Paragraph(
            f"<b>ORDER ID</b><br/>{html.escape(str(order['order_id']))}<br/>"
            f"<font color='#bfc7ff'>{html.escape(str(order['date']))}</font>",
            ParagraphStyle("OrderBox", parent=white_small, alignment=TA_RIGHT)
        )
    ]], colWidths=[320, 210])
    header.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#211d4f")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 18),
        ("RIGHTPADDING", (0,0), (-1,-1), 18),
        ("TOPPADDING", (0,0), (-1,-1), 17),
        ("BOTTOMPADDING", (0,0), (-1,-1), 17),
    ]))
    story += [header, Spacer(1, 14),
              Paragraph("Sales Invoice", section),
              Paragraph("Order fulfilment record generated by BooksKart.", small),
              Spacer(1, 12)]

    customer_table = Table([[
        Paragraph("<b>Customer</b><br/>" + html.escape(str(order["customer_name"])) +
                  "<br/>" + html.escape(str(order["phone"])), normal),
        Paragraph("<b>Delivery Address</b><br/>" + html.escape(address), normal)
    ]], colWidths=[265, 265])
    customer_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#f7f8fb")),
        ("BOX", (0,0), (-1,-1), 0.7, colors.HexColor("#e1e5ec")),
        ("INNERGRID", (0,0), (-1,-1), 0.5, colors.HexColor("#e5e8ee")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 12),
        ("RIGHTPADDING", (0,0), (-1,-1), 12),
        ("TOPPADDING", (0,0), (-1,-1), 11),
        ("BOTTOMPADDING", (0,0), (-1,-1), 11),
    ]))
    story += [customer_table, Spacer(1, 16)]

    rows = [[
        Paragraph("<b>Book</b>", normal),
        Paragraph("<b>Code</b>", normal),
        Paragraph("<b>Qty</b>", normal),
        Paragraph("<b>Unit Price</b>", normal),
        Paragraph("<b>Line Total</b>", normal),
    ]]
    for i in items:
        rows.append([
            Paragraph("◆ " + html.escape(str(i["name"])) + "<br/><font size='7' color='#667085'>" +
                      html.escape(str(i.get("author", ""))) + "</font>", normal),
            Paragraph(html.escape(str(i["book_id"])), small),
            Paragraph(str(i["quantity"]), normal),
            Paragraph(f"INR {i['unit_price']:,.2f}", normal),
            Paragraph(f"<b>INR {i['line_total']:,.2f}</b>", normal),
        ])
    table = Table(rows, colWidths=[220, 75, 45, 85, 85], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#eeeaff")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.HexColor("#342a79")),
        ("GRID", (0,0), (-1,-1), 0.45, colors.HexColor("#dfe3e9")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#fafbfc")]),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 7),
        ("RIGHTPADDING", (0,0), (-1,-1), 7),
        ("TOPPADDING", (0,0), (-1,-1), 8),
        ("BOTTOMPADDING", (0,0), (-1,-1), 8),
    ]))
    story += [table, Spacer(1, 15)]

    summary = Table([[
        Paragraph(
            "<b>Order Summary</b><br/>"
            f"Items: {sum(int(i['quantity']) for i in items)}<br/>"
            f"Payment: {html.escape(str(order['payment']))}<br/>"
            "<font color='#087f70'><b>Order Status: Confirmed</b></font>",
            normal
        ),
        Paragraph(f"<font size='9' color='#667085'>TOTAL</font><br/><b>INR {order['grand_total']:,.2f}</b>", total_style)
    ]], colWidths=[265, 265])
    summary.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#f7f8fb")),
        ("BOX", (0,0), (-1,-1), 0.7, colors.HexColor("#dfe3e9")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 13),
        ("RIGHTPADDING", (0,0), (-1,-1), 13),
        ("TOPPADDING", (0,0), (-1,-1), 12),
        ("BOTTOMPADDING", (0,0), (-1,-1), 12),
    ]))
    story += [summary, Spacer(1, 15),
              Paragraph("◆ BooksKart uses this document as the customer-facing order and payment record. Cash on Delivery is recorded as the selected payment method.", small),
              Spacer(1, 8),
              Paragraph("BooksKart · Designed & Developed By MJJ-TechWorld · Made In India", footer_style)]

    def footer(canvas, doc_obj):
        canvas.saveState()
        width, height = A4
        canvas.setStrokeColor(colors.HexColor("#d9dce5"))
        canvas.line(32, 27, width - 32, 27)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#69718c"))
        canvas.drawCentredString(width / 2, 16, "© 2026 BooksKart • Terms Of Service • Privacy Policy • Disclaimer • Version 2.0")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()

